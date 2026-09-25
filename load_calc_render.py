# load_calc_render.py
"""
Fill load_calculation.xlsx (C7 block) and render it to a PNG.

Merged-safe: writes to the top-left anchor of any merged range.

Renderer: Playwright + Chromium.
(wkhtmltopdf is no longer used — it is not available on Debian trixie,
the current Streamlit Cloud base image.)

Proposed load: always a single Nokia MF-2 OLT (fixed in load_calc_helper).
"""
from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import load_workbook

from load_calc_helper import compute_sufficiency, NOKIA_MF2_LOAD


SHEET_BY_INDEX = {
    1: "RS1-computation",
    2: "RS2-computation",
    3: "RS3-computation (existing)",
    4: "RS4-computation (existing)",
}


# ─────────────────────────────────────────────────────────────
# Merged-safe cell writer
# ─────────────────────────────────────────────────────────────

def _safe_set(ws, row: int, col: int, value):
    """
    Write `value` at (row, col). If the target is a MergedCell,
    walk to the top-left anchor of its merged range and write there.
    """
    cell = ws.cell(row=row, column=col)
    if type(cell).__name__ != "MergedCell":
        cell.value = value
        return

    for mr in ws.merged_cells.ranges:
        if (mr.min_row <= row <= mr.max_row) and \
           (mr.min_col <= col <= mr.max_col):
            ws.cell(row=mr.min_row, column=mr.min_col).value = value
            return

    try:
        cell.value = value
    except AttributeError:
        pass


# ─────────────────────────────────────────────────────────────
# Template filler
# ─────────────────────────────────────────────────────────────

def fill_template(
    template_path: str,
    out_path: str,
    data: dict,
    sheet_name: str,
) -> str:
    """Copy template → out_path, write values into the C7 block."""
    shutil.copy(template_path, out_path)
    wb = load_workbook(out_path)

    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found in workbook. "
            f"Available: {wb.sheetnames}"
        )
    ws = wb[sheet_name]

    # ---- Find label cells ----
    label_to_cell = _find_label_cells(ws)

    def put(label: str, value, *, row_offset: int = 0, col_offset: int = 8):
        """Write `value` `col_offset` columns right of the label."""
        cell = label_to_cell.get(label)
        if not cell:
            return
        _safe_set(ws, cell.row + row_offset, cell.column + col_offset, value)

    # ---- C7.1 (existing rectifier) ----
    put("1) EXISTING RECTIFIER SYSTEM BRAND - MODEL",
        data["rectifier_brand"])
    put("2) MAXIMUM RECTIFIER MODULES / INSTALLED",
        data["max_modules"])
    put("3) RECTIFIER MODULE RATING, WATTS / AMPS",
        data["module_rating_w"])
    put("4) BATTERY BRAND / VOLTAGE RATING",
        f'{data["battery_brand"]} {data["battery_voltage"]}'.strip())
    put("5) NUMBER OF BATTERIES in BANKS / CAPACITY per CELL (AH)",
        f'{data["battery_banks"]} / {data["battery_capacity_ah"]}')
    put("6) PRESENT LOAD READING, PANEL DISPLAY / CLAMP METER (A)",
        data["present_load_a"])
    put("7) RECTIFIER MODULES IN OPERATION",
        data["modules_in_operation"])
    put("8) ACTUAL FLOAT VOLTAGE (V)",
        data["actual_float_voltage_v"])

    # ---- C7.2 (proposed loads) — fixed Nokia MF-2, merged-safe ----
    start_row = _find_first_load_row(ws)

    # Write the single Nokia MF-2 row
    _safe_set(ws, start_row, 1,  NOKIA_MF2_LOAD["load_no"])
    _safe_set(ws, start_row, 3,  NOKIA_MF2_LOAD["equipment"])
    _safe_set(ws, start_row, 12, NOKIA_MF2_LOAD["power_w"])
    _safe_set(ws, start_row, 15, NOKIA_MF2_LOAD["current_a"])
    _safe_set(ws, start_row, 21, NOKIA_MF2_LOAD["cable_awg"])
    _safe_set(ws, start_row, 27, NOKIA_MF2_LOAD["breaker_a"])

    # Clear rows 2–7 so no stale template data lingers
    for r in range(start_row + 1, start_row + 7):
        for col in (1, 3, 12, 15, 21, 27):
            _safe_set(ws, r, col, "")

    # ---- Computed block ----
    comp = compute_sufficiency(data)

    _set_by_label(ws, label_to_cell, "Total Full Load Current",
                  comp["total_full_load_a"])
    _set_by_label(ws, label_to_cell, "Existing Rectifier Capacity:",
                  comp["existing_rectifier_capacity_a"])
    _set_by_label(ws, label_to_cell, "Existing Battery Capacity:",
                  comp["existing_battery_capacity_ah"])
    _set_by_label(ws, label_to_cell, "Available Rectifier Capacity =",
                  comp["available_rectifier_capacity_a"])
    _set_by_label(ws, label_to_cell, "Percent Utilization",
                  comp["percent_utilization"])
    _set_by_label(ws, label_to_cell, "BBUT =",
                  comp["bbut_hours"])

    wb.save(out_path)
    return out_path


# ─────────────────────────────────────────────────────────────
# Helpers
# ─────────────────────────────────────────────────────────────

def _find_label_cells(ws) -> dict:
    found = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                k = cell.value.strip()
                if k and k not in found:
                    found[k] = cell
    return found


def _find_first_load_row(ws) -> int:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == "Load No":
                return cell.row + 1
    print("⚠️ 'Load No' header not found — defaulting to row 21.")
    return 21


def _set_by_label(ws, label_cells: dict, prefix: str, value):
    for text, cell in label_cells.items():
        if text.startswith(prefix):
            _safe_set(ws, cell.row, cell.column + 1, value)
            return


# ─────────────────────────────────────────────────────────────
# PNG rendering — Playwright only
# ─────────────────────────────────────────────────────────────

def render_sheet_png(xlsx_path: str, sheet_name: str, out_png: str) -> str:
    """
    Render an XLSX sheet to a PNG using xlsx2html + Playwright.

    wkhtmltopdf is NOT used — it is no longer available on Debian
    trixie. Playwright + Chromium is the only renderer.
    """
    xlsx_path = str(xlsx_path)
    out_png = Path(out_png)
    out_png.parent.mkdir(parents=True, exist_ok=True)

    # 1) Build the HTML from the sheet
    try:
        from xlsx2html import xlsx2html
    except ImportError as e:
        raise RuntimeError(
            "Missing dependency: xlsx2html. Add it to requirements.txt "
            "and reboot the app."
        ) from e

    html_path = out_png.with_suffix(".html")
    with open(html_path, "w", encoding="utf-8") as f:
        # NOTE: xlsx2html uses `sheet=`, NOT `sheet_name=`
        xlsx2html(
            xlsx_path,
            sheet=sheet_name,
            output=f,
            locale="en_US",
        )

    # 2) Screenshot with Playwright
    try:
        from playwright.sync_api import sync_playwright
    except ImportError as e:
        raise RuntimeError(
            "Playwright not installed. Add 'playwright' to requirements.txt "
            "and create setup.sh with 'python -m playwright install chromium'."
        ) from e

    with sync_playwright() as p:
        browser = p.chromium.launch(args=["--no-sandbox"])
        page = browser.new_page(viewport={"width": 1400, "height": 1300})
        page.goto(f"file://{html_path.resolve()}")
        page.wait_for_timeout(600)
        page.screenshot(
            path=str(out_png),
            full_page=False,
            clip={"x": 0, "y": 0, "width": 1400, "height": 1300},
        )
        browser.close()

    if not out_png.exists() or out_png.stat().st_size < 1000:
        raise RuntimeError("Playwright did not produce a valid PNG.")

    return str(out_png)
